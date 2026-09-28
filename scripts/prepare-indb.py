#!/usr/bin/env python3
"""Audit a local INDB workbook; optionally export a private, opt-in catalog.

Install openpyxl separately to run this tool. No workbook data is shipped with
Mealio. Export requires explicit rights confirmation and stays in dataset/.
"""

import argparse
import json
import math
import os
import re
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DATASET = ROOT / "dataset"
DEFAULT_WORKBOOK = DATASET / "Anuvaad_INDB_2024.11.xlsx"
FIELDS = {
    "sourceId": "food_code", "name": "food_name",
    "kcalPer100g": "energy_kcal", "proteinPer100g": "protein_g",
    "carbsPer100g": "carb_g", "servingUnit": "servings_unit",
    "servingKcal": "unit_serving_energy_kcal",
    "servingProtein": "unit_serving_protein_g",
    "servingCarbs": "unit_serving_carb_g",
}
BASE = ("kcalPer100g", "proteinPer100g", "carbsPer100g")
SERVING = ("servingKcal", "servingProtein", "servingCarbs")


def numeric(value):
    if isinstance(value, bool) or not isinstance(value, (float, int)):
        return None
    return float(value) if math.isfinite(value) and value >= 0 else None


def audit_rows(rows):
    """Return validated records and aggregate findings, without changing rows."""
    records = []
    seen = set()
    report = {key: 0 for key in (
        "rows", "valid", "missing_name_or_code", "duplicate_codes",
        "invalid_nutrients", "missing_servings", "partial_servings",
        "unusual_serving_weights", "ambiguous_dish_aliases",
    )}
    for row in rows:
        report["rows"] += 1
        code, name = row.get("sourceId"), row.get("name")
        if not isinstance(code, str) or not code.strip() or not isinstance(name, str) or not name.strip():
            report["missing_name_or_code"] += 1
            continue
        code, name = code.strip(), name.strip()
        if code in seen:
            report["duplicate_codes"] += 1
            continue
        seen.add(code)
        base = {key: numeric(row.get(key)) for key in BASE}
        if any(value is None for value in base.values()):
            report["invalid_nutrients"] += 1
            continue

        unit = row.get("servingUnit")
        unit = unit.strip() if isinstance(unit, str) else ""
        values = {key: numeric(row.get(key)) for key in SERVING}
        missing = all(row.get(key) is None for key in SERVING)
        if not unit and missing:
            report["missing_servings"] += 1
            values = dict.fromkeys(SERVING)
        elif not unit or any(value is None for value in values.values()):
            report["partial_servings"] += 1
            # Retain reliable per-100g values, but disable this serving entirely.
            unit = ""
            values = dict.fromkeys(SERVING)
        elif base["kcalPer100g"] > 0:
            implied = 100 * values["servingKcal"] / base["kcalPer100g"]
            if implied < 5 or implied > 500:
                report["unusual_serving_weights"] += 1
        records.append({"sourceId": code, "name": name, **base, "servingUnit": unit or None, **values})
    report["valid"] = len(records)
    aliases = defaultdict(set)
    for record in records:
        for name in record["name"].split("(")[0].split("/"):
            key = re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]+", " ", name.lower())).strip()
            if key:
                aliases[key].add(record["sourceId"])
    report["ambiguous_dish_aliases"] = sum(len(codes) > 1 for codes in aliases.values())
    return records, report


def load_rows(path):
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise ValueError("Install openpyxl to audit the workbook.") from error
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        iterator = worksheet.iter_rows(values_only=True)
        headers = next(iterator)
        if not set(FIELDS.values()).issubset(set(headers)):
            raise ValueError("INDB workbook columns do not match the expected layout.")
        indices = {field: headers.index(column) for field, column in FIELDS.items()}
        return [
            {field: values[index] for field, index in indices.items()}
            for values in iterator if any(value is not None for value in values)
        ]
    finally:
        workbook.close()


def export_path(value):
    path = Path(value)
    path = (path if path.is_absolute() else ROOT / path).resolve()
    if path.suffix != ".json" or DATASET.resolve() not in path.parents:
        raise ValueError("JSON export must stay inside the ignored dataset/ directory.")
    return path


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--export", help="Optional private JSON path within dataset/")
    parser.add_argument("--confirm-rights", action="store_true", help="Confirm permission to use INDB data")
    args = parser.parse_args(argv)
    if args.export and not args.confirm_rights:
        parser.error("--export requires --confirm-rights after verifying dataset terms")
    if args.confirm_rights and not args.export:
        parser.error("--confirm-rights is only valid with --export")
    workbook = args.workbook if args.workbook.is_absolute() else ROOT / args.workbook
    records, report = audit_rows(load_rows(workbook))
    print(json.dumps(report, indent=2))
    if args.export:
        if report["valid"] != report["rows"]:
            parser.error("Export refused: missing/duplicate IDs or invalid base nutrients")
        path = export_path(args.export)
        path.parent.mkdir(parents=True, exist_ok=True)
        # No overwrite and owner-only permissions for a locally licensed copy.
        with os.fdopen(os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600), "w", encoding="utf8") as output:
            json.dump(records, output, ensure_ascii=False)
        print("Private catalog exported inside ignored dataset/.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
